import math2docx
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import os
import sys

# ---------- Regex ----------
INLINE_TOKEN_RE = re.compile(
    r'\*\*([\s\S]+?)\*\*'                                  # 1: bold
    r'|(?<!\*)\*(?!\*)([\s\S]+?)(?<!\*)\*(?!\*)'           # 2: italic
    r'|`([^`]+?)`'                                         # 3: code / `latex`
    r'|\\\(([\s\S]+?)\\\)'                                 # 4: \( ... \)
    r'|(?<!\$)\$(?!\$)([^\$\n]+?)\$(?!\$)'                 # 5: $ ... $
)

MATH_ENVS = ['equation', 'equation*', 'align', 'align*', 'gather', 'gather*',
             'eqnarray', 'eqnarray*', 'multline', 'multline*', 'displaymath',
             'math', 'split', 'aligned', 'cases', 'matrix', 'pmatrix',
             'bmatrix', 'vmatrix', 'array']
LIST_ENVS = ['itemize', 'enumerate', 'description']

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

# ---------- Utilitas LaTeX ----------
def preprocess_escaped_braces(s):
    """
    Ubah pola \\{...\\} (kurung kurawal ter-escape) menjadi
    \\lbrace ... \\rbrace agar math2docx tidak salah parse sebagai
    group LaTeX. Hanya menyentuh pasangan seimbang tanpa nested.
    """
    # Beberapa iterasi untuk menangani pasangan bersarang sederhana
    pattern = re.compile(r'\\\{([^{}]*?)\\\}')
    for _ in range(5):
        new_s = pattern.sub(r'\\lbrace \1 \\rbrace', s)
        if new_s == s:
            break
        s = new_s
    return s

def normalize_sqrt(s):
    s = re.sub(r'\\sqrt\[2\]\{([^}]*)\}', r'\\sqrt{\1}', s)
    s = re.sub(r'(\\sqrt)([^{])', r'\1{\2}', s)
    return s

def clean_math(s):
    s = preprocess_escaped_braces(s)
    s = normalize_sqrt(s)
    return s.strip()

def looks_like_latex(s):
    if not s:
        return False
    if re.search(r'\\[a-zA-Z]+', s):
        return True
    if re.search(r'[\^_]\{', s):
        return True
    return False

def strip_unpaired_markers(text):
    if text.count('**') % 2 == 1:
        idx = text.rfind('**')
        if idx != -1:
            text = text[:idx] + text[idx+2:]
    return text

# ---------- Paksa warna hitam ----------
def force_black_runs(paragraph):
    for elem in paragraph._element.iter():
        if not elem.tag.endswith('}r'):
            continue
        rPr = elem.find(f'{{{W_NS}}}rPr')
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            elem.insert(0, rPr)
        for c in list(rPr.findall(f'{{{W_NS}}}color')):
            rPr.remove(c)
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000')
        rPr.append(color)

# ---------- Sembunyikan placeholder derajat akar ----------
def hide_radical_degree(paragraph):
    for rad in paragraph._element.iter(f'{{{M_NS}}}rad'):
        radPr = rad.find(f'{{{M_NS}}}radPr')
        if radPr is None:
            radPr = OxmlElement('m:radPr')
            rad.insert(0, radPr)
        if radPr.find(f'{{{M_NS}}}degHide') is None:
            degHide = OxmlElement('m:degHide')
            degHide.set(qn('m:val'), '1')
            radPr.append(degHide)

# ---------- Sembunyikan placeholder batas sum/integral yang kosong ----------
def hide_empty_nary_limits(paragraph):
    """
    Sembunyikan placeholder batas atas/bawah n-ary yang kosong.
    Deteksi 'kosong' juga mencakup run tanpa teks.
    """
    for nary in paragraph._element.iter(f'{{{M_NS}}}nary'):
        sup = nary.find(f'{{{M_NS}}}sup')
        sub = nary.find(f'{{{M_NS}}}sub')

        def is_empty(lim):
            if lim is None:
                return True
            for t in lim.iter(f'{{{M_NS}}}t'):
                if t.text and t.text.strip():
                    return False
            return True

        sup_empty = is_empty(sup)
        sub_empty = is_empty(sub)

        if not (sup_empty or sub_empty):
            continue

        naryPr = nary.find(f'{{{M_NS}}}naryPr')
        if naryPr is None:
            naryPr = OxmlElement('m:naryPr')
            nary.insert(0, naryPr)

        chr_elem    = naryPr.find(f'{{{M_NS}}}chr')
        limLoc_elem = naryPr.find(f'{{{M_NS}}}limLoc')
        grow_elem   = naryPr.find(f'{{{M_NS}}}grow')

        for c in list(naryPr):
            naryPr.remove(c)

        if chr_elem is not None:    naryPr.append(chr_elem)
        if limLoc_elem is not None: naryPr.append(limLoc_elem)
        if grow_elem is not None:   naryPr.append(grow_elem)

        if sub_empty:
            subHide = OxmlElement('m:subHide')
            subHide.set(qn('m:val'), '1')
            naryPr.append(subHide)

        if sup_empty:
            supHide = OxmlElement('m:supHide')
            supHide.set(qn('m:val'), '1')
            naryPr.append(supHide)

# ---------- Helper: bersihkan semua math di paragraf ----------
def fix_math_in_paragraph(paragraph):
    hide_radical_degree(paragraph)
    hide_empty_nary_limits(paragraph)

# ---------- Smart split untuk tabel ----------
def split_table_row(row):
    row = row.strip()
    if row.startswith('|'):
        row = row[1:]
    if row.endswith('|'):
        row = row[:-1]

    cells, current = [], []
    i = 0
    depth_brace = 0
    in_dollar = False
    in_inline_paren = False

    while i < len(row):
        ch = row[i]
        if ch == '\\' and i + 1 < len(row):
            nxt = row[i + 1]
            if nxt == '(':
                in_inline_paren = True
            elif nxt == ')':
                in_inline_paren = False
            current.append(ch)
            current.append(nxt)
            i += 2
            continue
        if ch == '$':
            in_dollar = not in_dollar
            current.append(ch)
        elif ch == '{':
            depth_brace += 1
            current.append(ch)
        elif ch == '}':
            depth_brace -= 1
            current.append(ch)
        elif ch == '|' and not in_dollar and not in_inline_paren and depth_brace == 0:
            cells.append(''.join(current).strip())
            current = []
        else:
            current.append(ch)
        i += 1
    cells.append(''.join(current).strip())
    return cells

# ---------- Render inline ----------
def add_inline_content(paragraph, text, bold=False, italic=False):
    if not bold and not italic:
        text = strip_unpaired_markers(text)

    pos = 0
    for m in INLINE_TOKEN_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.bold = bold
            r.italic = italic

        if m.group(1) is not None:
            add_inline_content(paragraph, m.group(1), bold=True, italic=italic)

        elif m.group(2) is not None:
            add_inline_content(paragraph, m.group(2), bold=bold, italic=True)

        elif m.group(3) is not None:
            code = m.group(3)
            if looks_like_latex(code):
                try:
                    math2docx.add_math(paragraph, clean_math(code))
                except Exception:
                    r = paragraph.add_run(code)
                    r.font.name = 'Consolas'
                    r.bold = bold; r.italic = italic
            else:
                r = paragraph.add_run(code)
                r.font.name = 'Consolas'
                r.bold = bold; r.italic = italic

        elif m.group(4) is not None or m.group(5) is not None:
            latex = m.group(4) if m.group(4) is not None else m.group(5)
            try:
                math2docx.add_math(paragraph, clean_math(latex))
            except Exception:
                r = paragraph.add_run(latex)
                r.bold = bold; r.italic = italic

        pos = m.end()

    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = bold
        r.italic = italic

def detect_begin_env(line):
    m = re.match(r'^\\begin\{([^}]+)\}', line)
    return m.group(1) if m else None

# ---------- Parser ----------
def parse_content(content, debug=False):
    lines = content.splitlines()
    elements = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        env = detect_begin_env(stripped)
        if env:
            end_marker = f'\\end{{{env}}}'
            body_lines = []
            rest = stripped[len(f'\\begin{{{env}}}'):]
            if end_marker in rest:
                idx = rest.find(end_marker)
                body_lines.append(rest[:idx])
                i += 1
            else:
                if rest.strip():
                    body_lines.append(rest)
                i += 1
                while i < len(lines):
                    if end_marker in lines[i]:
                        idx = lines[i].find(end_marker)
                        body_lines.append(lines[i][:idx])
                        break
                    body_lines.append(lines[i])
                    i += 1
                i += 1
            body = '\n'.join(body_lines).strip()
            if env in MATH_ENVS:
                elements.append(("block_math", body))
            elif env in LIST_ENVS:
                for it in re.split(r'\\item\s*', body):
                    it = it.strip()
                    if it:
                        elements.append(("list_item", it))
            else:
                elements.append(("block_math", body))
                if debug:
                    print(f"   ⚠️  Environment tidak dikenal: {env}")
            continue

        is_bracket = stripped.startswith('\\[')
        is_dollar  = stripped.startswith('$$') and not stripped.startswith('$$$')
        if is_bracket or is_dollar:
            open_delim  = '\\[' if is_bracket else '$$'
            close_delim = '\\]' if is_bracket else '$$'
            rest = stripped[len(open_delim):]
            if close_delim in rest:
                idx = rest.find(close_delim)
                elements.append(("block_math", rest[:idx].strip()))
                i += 1
                continue
            math_lines = [rest] if rest else []
            i += 1
            while i < len(lines):
                l = lines[i]
                if close_delim in l:
                    idx = l.find(close_delim)
                    math_lines.append(l[:idx])
                    break
                math_lines.append(l)
                i += 1
            elements.append(("block_math", '\n'.join(math_lines).strip()))
            i += 1
            continue

        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            elements.append(("heading", level, stripped[level:].strip()))
            i += 1
            continue

        if re.match(r'^([-*_])\1{2,}$', stripped):
            elements.append(("hr",))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            elements.append(("table", table_lines))
            continue

        if re.match(r'^\s*[-*+]\s+', line) or re.match(r'^\s*\d+\.\s+', line):
            elements.append(("list_item", stripped))
            i += 1
            continue

        elements.append(("paragraph", stripped))
        i += 1

    return elements

# ---------- Render ----------
def render(doc, elements):
    for elem in elements:
        kind = elem[0]

        if kind == "heading":
            h = doc.add_heading("", level=min(elem[1], 4))
            add_inline_content(h, elem[2])
            fix_math_in_paragraph(h)
            force_black_runs(h)

        elif kind == "hr":
            doc.add_paragraph()

        elif kind == "paragraph":
            p = doc.add_paragraph()
            add_inline_content(p, elem[1])
            fix_math_in_paragraph(p)

        elif kind == "list_item":
            text = re.sub(r'^\s*[-*+]\s+', '', elem[1])
            text = re.sub(r'^\s*\d+\.\s+', '', text)
            p = doc.add_paragraph(style='List Bullet')
            add_inline_content(p, text)
            fix_math_in_paragraph(p)

        elif kind == "block_math":
            p = doc.add_paragraph()
            try:
                math2docx.add_math(p, clean_math(elem[1]))
                fix_math_in_paragraph(p)
            except Exception:
                p.add_run(f"[{elem[1]}]")

        elif kind == "table":
            rows = elem[1]
            cells = []
            for row in rows:
                if re.match(r'^\|[\s\-:|]+\|$', row):
                    continue
                cells.append(split_table_row(row))
            if cells:
                ncols = max(len(r) for r in cells)
                table = doc.add_table(rows=len(cells), cols=ncols)
                table.style = 'Table Grid'
                for i, row in enumerate(cells):
                    for j, cell_text in enumerate(row):
                        if j < ncols:
                            cell = table.cell(i, j)
                            cell.text = ""
                            add_inline_content(cell.paragraphs[0], cell_text)
                            fix_math_in_paragraph(cell.paragraphs[0])

# ---------- Main ----------
def main():
    args = sys.argv[1:]
    debug = False
    if '--debug' in args:
        debug = True
        args.remove('--debug')

    input_file  = args[0] if len(args) > 0 else "input.txt"
    output_file = args[1] if len(args) > 1 else "output.docx"

    if not os.path.isfile(input_file):
        if os.path.isfile("rumus.txt"):
            input_file = "rumus.txt"
        else:
            print(f"❌ File '{input_file}' tidak ditemukan.")
            sys.exit(1)

    with open(input_file, 'r', encoding='utf-8') as f:
        content = f.read()

    if not content.strip():
        print("❌ File kosong.")
        sys.exit(1)

    elements = parse_content(content, debug=debug)
    print(f"🔍 Ditemukan {len(elements)} elemen dari '{input_file}'.")

    if debug:
        from collections import Counter
        counts = Counter(e[0] for e in elements)
        for k, v in counts.items():
            print(f"   • {k}: {v}")

    doc = Document()
    render(doc, elements)

    try:
        doc.save(output_file)
        print(f"✅ Berhasil menyimpan {output_file}")
    except PermissionError:
        base, ext = os.path.splitext(output_file)
        new_name = f"{base}_baru{ext}"
        doc.save(new_name)
        print(f"⚠️ File '{output_file}' sedang terbuka. Disimpan sebagai '{new_name}'")
        output_file = new_name

    try:
        os.startfile(output_file)
        print("📂 File dibuka otomatis.")
    except Exception as e:
        print(f"⚠️ Tidak bisa membuka file: {e}")

if __name__ == "__main__":
    main()